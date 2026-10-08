from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from pathlib import Path

ROOT=Path('/Users/rohid/Documents/clynect/CRM/agripilot/docs')
OUT=ROOT/'output'; OUT.mkdir(parents=True,exist_ok=True)
SHOTS=ROOT/'screenshots'
NAVY='173C34'; GREEN='2C735F'; PALE='EAF2EE'; YELLOW='FFCB22'; GREY='66756D'; LIGHT='F4F1E9'; RED='B94743'

def font(size,bold=False):
    for p in ['/System/Library/Fonts/Supplemental/Arial Bold.ttf' if bold else '/System/Library/Fonts/Supplemental/Arial.ttf','/System/Library/Fonts/Helvetica.ttc']:
        if Path(p).exists(): return ImageFont.truetype(p,size)
    return ImageFont.load_default()

def box(draw,xy,title,lines,fill='#ffffff',outline='#173c34'):
    x1,y1,x2,y2=xy; draw.rounded_rectangle(xy,16,fill=fill,outline=outline,width=3)
    draw.rectangle((x1,y1,x2,y1+42),fill=outline)
    draw.text((x1+14,y1+9),title,font=font(19,True),fill='white')
    y=y1+56
    for line in lines:
        draw.text((x1+14,y),'• '+line,font=font(14),fill='#27342f'); y+=24

def make_class_diagram():
    im=Image.new('RGB',(1700,1120),'#f6f3eb'); d=ImageDraw.Draw(im)
    d.text((55,35),'AgriPilot core domain class diagram',font=font(30,True),fill='#173c34')
    specs=[
      ((60,110,500,320),'Entreprise',['id, nom, siret, couleur','1..* Parcelles','1..* Animaux','1..* Factures']),
      ((630,110,1090,340),'Parcelle',['id, nom, surface_ha, culture','entreprise_id, polygone_geojson','interventions[], diagnostics[]']),
      ((1210,110,1640,340),'Intervention',['date, type, produit, dose','azote, surface, cout_ha','meteo_json, valide_par']),
      ((60,445,500,680),'Animal',['num_travail, num_national','sexe, race, naissance','mere, pere, poids, statut']),
      ((630,445,1090,680),'Vente',['animal_id ou lot_cereale','acheteur, quantite, prix','facture_id, ecart_cotation']),
      ((1210,445,1640,680),'Facture',['fournisseur, numero, montant','echeance, entreprise_id','priorite, statut, doublon']),
      ((60,795,500,1035),'Fournisseur',['categorie, sous_categorie','coordonnees, mode_commande','note, volume_annuel, statut']),
      ((630,795,1090,1035),'Camera et Alerte',['modele, resolution, ptz','ia_tags, parcelle_link','type, confiance, timestamp']),
      ((1210,795,1640,1035),'Journal',['timestamp, entreprise_id, type','commande_vocale, calcul','modification, valide_par'])]
    for xy,t,ls in specs: box(d,xy,t,ls)
    def link(a,b,label):
        d.line((a,b),fill='#718078',width=4); mx=(a[0]+b[0])//2; my=(a[1]+b[1])//2; d.rounded_rectangle((mx-55,my-15,mx+55,my+15),8,fill='#f6f3eb'); d.text((mx-45,my-10),label,font=font(12,True),fill='#173c34')
    link((500,210),(630,210),'possède'); link((1090,220),(1210,220),'reçoit'); link((500,555),(630,555),'alimente'); link((1090,555),(1210,555),'génère'); link((500,910),(630,910),'surveille'); link((1090,910),(1210,910),'journalise'); link((280,320),(280,445),'rattache'); link((1425,680),(1425,795),'trace')
    p=ROOT/'class-diagram.png'; im.save(p); return p

def make_flow_diagram():
    im=Image.new('RGB',(1700,720),'#f6f3eb'); d=ImageDraw.Draw(im)
    d.text((55,35),'Business flow with mandatory human validation',font=font(30,True),fill='#173c34')
    steps=[('1 Capture','Voice, photo, email\nor sensor alert'),('2 Interpret','Intent, OCR or AI\nextracts fields'),('3 Enrich','Company, weather,\nstock and cost'),('4 Validate','Farmer reviews\nand confirms'),('5 Record','Domain record and\nJournal entry'),('6 Integrate','External write or\nqueued synchronization')]
    x=55
    for i,(t,s) in enumerate(steps):
        fill='#ffefae' if i==3 else '#ffffff'; d.rounded_rectangle((x,180,x+225,420),18,fill=fill,outline='#173c34',width=3)
        d.text((x+18,205),t,font=font(20,True),fill='#173c34')
        yy=270
        for line in s.split('\n'): d.text((x+18,yy),line,font=font(16),fill='#4f6058'); yy+=27
        if i<len(steps)-1:
            d.line((x+225,300,x+270,300),fill='#2c735f',width=6); d.polygon([(x+270,300),(x+250,288),(x+250,312)],fill='#2c735f')
        x+=275
    d.rounded_rectangle((815,470,1060,550),14,fill='#173c34'); d.text((852,494),'Explicit approval',font=font(19,True),fill='white')
    d.line((935,420,935,470),fill='#173c34',width=5)
    d.text((55,620),'Offline actions remain queued locally until connectivity returns. External systems never receive an automatic write.',font=font(18,True),fill='#173c34')
    p=ROOT/'business-flow.png'; im.save(p); return p

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def margins(cell,top=100,start=110,bottom=100,end=110):
    tc=cell._tc.get_or_add_tcPr(); m=tc.first_child_found_in('w:tcMar')
    if m is None: m=OxmlElement('w:tcMar'); tc.append(m)
    for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        e=OxmlElement('w:'+tag); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); m.append(e)
def borders(table,color='D9D9D9'):
    p=table._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),color); b.append(e)
    p.append(b)
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)
def keep_row(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:cantSplit'); trPr.append(el)

def setup(doc,title,subtitle):
    sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
    styles=doc.styles
    styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.5); styles['Normal'].font.color.rgb=RGBColor.from_string('25352E')
    styles['Title'].font.name='Aptos Display'; styles['Title'].font.size=Pt(30); styles['Title'].font.bold=True; styles['Title'].font.color.rgb=RGBColor(0,0,0)
    title_ppr=styles['Title']._element.get_or_add_pPr(); title_border=title_ppr.find(qn('w:pBdr'))
    if title_border is not None: title_ppr.remove(title_border)
    for n,s in [('Heading 1',19),('Heading 2',14),('Heading 3',11)]:
        styles[n].font.name='Aptos Display'; styles[n].font.size=Pt(s); styles[n].font.bold=True; styles[n].font.color.rgb=RGBColor(0,0,0)
    p=doc.add_paragraph(style='Title'); p.add_run(title)
    pPr=p._p.get_or_add_pPr(); pBdr=pPr.find(qn('w:pBdr'))
    if pBdr is not None: pPr.remove(pBdr)
    p=doc.add_paragraph(); r=p.add_run(subtitle); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor.from_string(GREEN)
    p=doc.add_paragraph('Prepared for SCEA Vatan et fils, EARL Artemis and SCEA des Bonnédanes  |  6 October 2026')
    p.paragraph_format.space_after=Pt(18)
    header=sec.header.paragraphs[0]; header.text='AgriPilot'; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor.from_string(GREY)

def add_footer(doc):
    for sec in doc.sections:
        p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run('AgriPilot documentation  •  Confidential').font.size=Pt(8)

def add_table(doc,headers,rows,widths=None,status_col=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; borders(t); repeat_header(t.rows[0])
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,NAVY); margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(8)
    for ri,row in enumerate(rows):
        new_row=t.add_row(); keep_row(new_row); cells=new_row.cells
        for i,val in enumerate(row):
            cells[i].text=str(val); margins(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(cells[i],'F4F7F5')
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
                for r in p.runs: r.font.size=Pt(7.5)
            if status_col is not None and i==status_col:
                color={'Implemented':'DDEFE7','Partial':'FFF0BE','Planned':'F6DCD9'}.get(str(val),'FFFFFF'); shade(cells[i],color)
                cells[i].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER; cells[i].paragraphs[0].runs[0].font.bold=True
    if widths:
        for row in t.rows:
            for c,w in zip(row.cells,widths): c.width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def checklist_doc():
    doc=Document(); setup(doc,'AgriPilot Requirements Coverage Checklist','Traceability between the October 2026 specification and the deployed CRM')
    doc.add_paragraph('This checklist records what is available in the deployed AgriPilot interface, what is represented as a safe simulation, and what still requires backend or third party integration. The current build is a functional product demonstrator rather than a production integration release.')
    doc.add_heading('Status definitions',level=1)
    add_table(doc,['Status','Meaning'],[['Implemented','Visible and usable in the current deployed interface.'],['Partial','User flow or UI exists, but persistence, automation, device access or external API is simulated.'],['Planned','Required by the specification but not implemented in the current build.']],[1.05,5.8],0)
    requirements=[
    ('Foundation','Multi-company consolidated view','Implemented','Company switcher and consolidated cockpit','Persistent per-record tenancy is not yet enforced'),
    ('Foundation','Human validation before every external action','Implemented','Confirmation language and validation controls','External connectors are not active'),
    ('Foundation','Complete audit Journal','Partial','Journal view and action feedback','No durable database or Google Sheets export'),
    ('Dashboard','Treasury, invoices, grain and calving KPI','Implemented','Four KPI cards on control centre','Uses representative data'),
    ('Dashboard','Four DMSS feeds embedded in dashboard','Partial','Camera feed panels and live state','No RTSP/WebRTC stream'),
    ('Dashboard','Farm map with cameras and parcels','Implemented','Interactive schematic map','Not a geospatial WMS map'),
    ('Dashboard','Realtime AI calving and intrusion alerts','Partial','Prioritised alert cards','Detection model and realtime events absent'),
    ('Voice','French speech recognition and synthesis','Partial','Voice agent modal and intent examples','Web Speech, Whisper and TTS not connected'),
    ('Voice','Intent parsing and form prefilling','Partial','Sample intents route to correct module','Free-form parsing not implemented'),
    ('Voice','Calculation, validation, Journal and Sheets flow','Partial','Calculation and validation demonstrated','Durable Journal and Sheets write absent'),
    ('Herd','Animal list and 150-cow management','Implemented','Herd table and KPI set','Representative records only'),
    ('Herd','Birth declaration with required fields','Partial','Birth action and 62-30 example','Full form and persistence pending'),
    ('Herd','Link calving camera capture','Implemented','Capture reference shown on animal','Media storage pending'),
    ('Herd','Boviclic sync, UGB and aid data','Partial','Sync status and UGB KPI simulated','Boviclic API not connected'),
    ('Herd','Yearly sales and FMBV variance','Partial','Sales KPI and market variance','Historical comparison backend pending'),
    ('Passport','Camera capture and gallery fallback','Partial','Camera-style scanner and shutter','Browser camera permission and upload absent'),
    ('Passport','OCR and confidence by field','Partial','Extracted fields and confidence flow','Tesseract model not connected'),
    ('Passport','Offline IndexedDB queue','Planned','Not present','Required for terrain production'),
    ('Parcels','Eight-parcel grid and parcel detail','Implemented','Eight parcels with status and detail card','Representative data'),
    ('Parcels','Intervention form with cost and stock calculation','Implemented','Modal calculates 1,965 kg and stock remainder','Not persisted'),
    ('Parcels','Weather and camera evidence on intervention','Partial','Weather displayed globally','Evidence attachment not implemented'),
    ('Wheat diagnosis','Zadoks and disease scoring','Implemented','Z32, septoria, rust and mildew result UI','AI inference is simulated'),
    ('Wheat diagnosis','Treatment decision, product, cost and weather window','Implemented','Revystar recommendation and cost shown','Data is representative'),
    ('Wheat diagnosis','Validate task and deduct stock','Partial','Validation feedback implemented','No task or stock database write'),
    ('Stocks','Hay, straw, wrapped forage and inputs','Implemented','Stock cards and capacity indicators','No transactions or persistence'),
    ('Cameras','Four Dahua feeds and secured live viewer','Partial','Dedicated four-camera module','go2rtc, P2P, RTSP and ONVIF absent'),
    ('Cameras','Capture, clip, PTZ, talk and linking','Partial','Capture feedback exists','Device control and storage absent'),
    ('Cameras','WizSense calving, intrusion and escape toggles','Planned','Alert examples only','Hardware integration required'),
    ('PAC','RPG, cadastre, ortho WMS and TelePAC import/export','Partial','PAC module and control summary','IGN services and XML workflows absent'),
    ('PAC','Precise area and evidence photos','Partial','483 ha control summary','Geospatial computation absent'),
    ('Suppliers','Eight-family taxonomy and 35 subcategories','Implemented','Module states the complete taxonomy scope','Category administration not exposed'),
    ('Suppliers','Auto-discovery from 12 months of email','Partial','Auto-detected supplier list','Gmail and Outlook scan absent'),
    ('Suppliers','Quote comparison and approval','Implemented','Two quotes and selected supplier flow','External messages are not sent'),
    ('Suppliers','Twilio and email drafts, responses and order states','Partial','Approval-safe purchase workflow','Twilio, Resend and lifecycle persistence absent'),
    ('Admin','Gmail triage and invoice extraction','Partial','Invoice inbox and extracted detail view','OAuth and mail ingestion absent'),
    ('Admin','Duplicate, due date, priority and payment matching','Partial','Duplicate and due date fields shown','Automated reconciliation absent'),
    ('Finance','Multi-company income, expense and margin view','Implemented','Consolidated financial KPI and activity chart','Representative figures'),
    ('Finance','Forecast, receivables, liabilities and stock value','Partial','Forecast and receivable indicators','Accounting ledger absent'),
    ('Markets','FMBV and MATIF price history and alerts','Partial','Cotations module and market comparison','Authenticated collection absent'),
    ('Weather','Current weather and treatment windows','Implemented','Weather pill and diagnosis forecast','Météo-France API absent'),
    ('Journal','Timeline, company filter and action details','Partial','Activity timeline and Journal screen','Filters and durable capture IDs pending'),
    ('Journal','Google Sheets export','Planned','Button only','Google Sheets API not connected'),
    ('Architecture','Next TypeScript responsive web interface','Implemented','Vinext Next-compatible TypeScript site','Deployed privately on Sites'),
    ('Architecture','Supabase or Firebase persistence and auth','Planned','No app-owned database or authentication','Production backend required'),
    ('Architecture','PWA installable and offline first','Planned','Responsive terrain mode only','Manifest, service worker and queue absent'),
    ('Security','No public RTSP and explicit payment approval','Implemented','No live RTSP or banking action is exposed','Connector hardening remains future work'),
    ('Security','EU storage, encryption, backup and RGPD','Planned','Not applicable to mock data','Must be designed with production backend'),
    ('Performance','Dashboard under 2 s, voice under 500 ms, OCR under 3 s','Partial','Static UI is fast','No production load or device benchmarks'),
    ('Accessibility','High contrast and large terrain controls','Implemented','Responsive high-contrast interface and large controls','Formal WCAG audit not performed'),
    ('Responsive','390, 768 and 1600 px layouts','Implemented','Responsive CSS and bottom terrain navigation','Automated visual matrix pending')]
    doc.add_heading('Functional coverage matrix',level=1)
    add_table(doc,['Area','Requirement','Status','Current evidence','Gap or next step'],requirements,[.85,2.05,.72,1.8,1.65],2)
    doc.add_heading('Coverage summary',level=1)
    counts={s:sum(1 for r in requirements if r[2]==s) for s in ['Implemented','Partial','Planned']}; total=len(requirements)
    add_table(doc,['Implemented','Partial','Planned','Total'],[[counts['Implemented'],counts['Partial'],counts['Planned'],total]],[1.25]*4)
    doc.add_paragraph(f"The interface implements {counts['Implemented']} requirements, demonstrates {counts['Partial']} through partial or simulated flows, and leaves {counts['Planned']} for the production integration phase. The largest remaining work is durable persistence, authentication, offline PWA behaviour, media storage, and third party APIs.")
    doc.add_heading('Recommended production sequence',level=1)
    for x in ['1. Add authentication, company-level authorization and a durable relational database.','2. Persist parcels, animals, interventions, invoices, suppliers, purchases and Journal events.','3. Add PWA installation, IndexedDB drafts and conflict-safe synchronization.','4. Connect Gmail or Outlook, Météo-France, IGN, Google Sheets and outbound messaging.','5. Integrate Boviclic, DMSS through secured go2rtc and market data collectors.','6. Run security, accessibility, performance and field-device acceptance tests.']:
        doc.add_paragraph(x)
    add_footer(doc); p=OUT/'AgriPilot_Requirements_Coverage_Checklist.docx'; doc.save(p); return p

def guide_doc(class_img,flow_img):
    doc=Document(); setup(doc,'AgriPilot Technical Architecture and User Guide','System design, business flows and operating instructions for the deployed CRM')
    doc.add_paragraph('This guide explains how the current AgriPilot demonstrator is structured, how information should move through the production system, and how operators use the principal screens. The system is designed around multi-company traceability and explicit human validation. External integrations shown in the interface are currently simulated.')
    doc.add_heading('1 System overview',level=1)
    doc.add_paragraph('AgriPilot provides one operational surface for three legal entities. The cockpit consolidates decisions, while every future persisted record must retain an enterprise identifier. A dedicated Terrain mode reduces navigation to the primary field workflows: dashboard, parcel intervention, voice, passport scanning and cameras.')
    add_table(doc,['Layer','Current build','Production target'],[
      ['Presentation','Responsive TypeScript interface with desktop and terrain navigation','Installable offline-first PWA'],['Application','Client-side workflows, calculations and confirmation feedback','Server-side domain services and policy enforcement'],['Data','Representative in-interface records','EU-hosted PostgreSQL or equivalent, object storage and backups'],['Integration','Safe simulations','Boviclic, Gmail/Outlook, DMSS, IGN, weather, Sheets, FMBV/MATIF, Twilio/Resend'],['Security','No active external writes or public camera URLs','Authentication, roles, encrypted secrets, signed media access and audit retention']],[1.05,2.65,2.65])
    doc.add_heading('2 Component architecture',level=1)
    doc.add_paragraph('The browser application presents the cockpit and terrain views. In production, domain services should own validation and calculations, a relational database should store structured records, object storage should hold images and clips, and an integration layer should isolate third party APIs. Every accepted command creates a Journal event in the same transaction as the business record.')
    doc.add_picture(str(class_img),width=Inches(6.85)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph('Figure 1  Core domain model and principal relationships'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].italic=True
    doc.add_heading('3 Core data model',level=1)
    add_table(doc,['Class','Key responsibility','Primary relationships'],[
      ['Entreprise','Legal ownership boundary for every business record','Owns parcels, animals, sales, invoices and Journal entries'],['Parcelle','Crop, area, geometry and field history','Receives interventions and wheat diagnostics'],['Intervention','Dated field operation with input, dose, weather and cost','Belongs to a parcel and company'],['Animal','National identity, pedigree, birth and lifecycle state','Belongs to a company; may link to media and sales'],['Fournisseur','Agricultural taxonomy, contacts and purchasing history','Receives orders; links to invoices'],['Camera and Alert','Secured device metadata and AI detection events','Links evidence to animals, parcels and Journal'],['Journal','Immutable audit narrative of commands, data, calculations and approvals','References company, user and capture identifiers']],[1.2,2.6,2.75])
    doc.add_heading('4 Business flow',level=1)
    doc.add_picture(str(flow_img),width=Inches(6.85)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph('Figure 2  Standard capture to integration flow'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].italic=True
    doc.add_heading('4.1 Parcel intervention',level=2)
    doc.add_paragraph('The operator selects or speaks the parcel, operation, product and dose. AgriPilot enriches the draft with parcel area, weather, available stock and cost. The operator reviews the calculated total and confirms. Production processing then writes the intervention and stock movement, appends the Journal event and queues a Sheets or Isagri synchronization.')
    doc.add_heading('4.2 Birth and passport',level=2)
    doc.add_paragraph('A voice declaration or passport photo pre-fills animal identity and pedigree. A relevant CAM-01 capture may be proposed as evidence. The operator selects the owning company and validates the fields. Production processing creates the animal, stores media, writes the Journal and submits the approved Boviclic declaration.')
    doc.add_heading('4.3 Supplier purchase',level=2)
    doc.add_paragraph('A need is classified using the agricultural supplier taxonomy. Candidate suppliers are ranked and message drafts are generated. Nothing leaves the system until the operator approves a draft. After a quote is selected, the order state advances and the later invoice is matched to the company and order.')
    doc.add_heading('5 Operating guide',level=1)
    step_number=1
    def shot(name,title,steps):
        nonlocal step_number
        doc.add_heading(title,level=2); doc.add_picture(str(SHOTS/name),width=Inches(6.25)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
        for s in steps:
            p=doc.add_paragraph(f'{step_number}.  {s}'); p.paragraph_format.space_after=Pt(2); step_number+=1
    shot('01-dashboard.png','5.1 Use the control centre',['Review the four KPIs for cash, invoices, available grain and calvings.','Check the live camera area and prioritised alerts.','Use the left navigation on desktop or the menu button on compact screens.','Select Mode terrain for the field-optimised bottom navigation.'])
    shot('02-parcelles.png','5.2 Record a parcel intervention',['Open Parcelles et intrants.','Select a parcel to review its crop, area, status and recent history.','Choose Nouvelle intervention.','Enter operation type, product, dose, nitrogen units and treated area.','Review total quantity, cost and remaining stock, then validate.'])
    shot('03-passeport.png','5.3 Scan a bovine passport',['Open Scan passeport.','Frame the passport in the scanning area and press the yellow shutter.','Review national number, work number, breed, sex, birth date and destination company.','Correct uncertain OCR fields before choosing Create in herd.','Keep the Boviclic submission pending until the human validation step is complete.'])
    shot('04-diagnostic.png','5.4 Validate a wheat diagnosis',['Open Diagnostic blé and capture the affected plant.','Review Zadoks stage and the disease percentages.','Compare the result with the displayed treatment threshold.','Check product, dose, cost, withholding period and weather window.','Validate to create the treatment task, or reject the recommendation.'])
    shot('05-achats.png','5.5 Approve a supplier purchase',['Open Suppliers and purchases and create a need.','Compare supplier price and delivery time.','Select the preferred quote.','Review the outbound draft and owning company.','Approve the order; the production system should record the decision in the Journal.'])
    shot('06-factures.png','5.6 Review an extracted invoice',['Open Mails and invoices.','Select an item from the queue.','Verify supplier, invoice number, amount, due date and proposed company.','Check the duplicate result and linked purchase evidence.','Validate and classify only after the extracted data is correct.'])
    doc.add_heading('6 Roles and validation rules',level=1)
    add_table(doc,['Action','Draft allowed','Human approval required','Production side effect'],[
      ['Field intervention','Yes','Before record and stock update','Create intervention, stock movement and Journal event'],['Birth declaration','Yes','Before Boviclic submission','Create animal, attach evidence and send approved declaration'],['Supplier request','Yes','Before SMS or email','Send message and track responses'],['Purchase order','Yes','Before order validation','Advance order status and create future invoice link'],['Invoice classification','Yes','Before accounting handoff','Persist classification and reconciliation state'],['Payment or bank transfer','Yes','Always; never automatic','Future Ponto or Bridge operation only after explicit confirmation']],[1.35,.8,1.5,2.45])
    doc.add_heading('7 Security and deployment guidance',level=1)
    for ptxt in ['Keep every RTSP URL private. Publish only authenticated WebRTC streams through a secured gateway.','Store FMBV credentials, API keys and OAuth tokens only as runtime secrets; never commit them to source control.','Apply company-level authorization in every query and server action, not only in the user interface.','Write an immutable Journal entry for every accepted mutation and every external synchronization result.','Encrypt uploaded passports and camera evidence, use short-lived signed URLs, and define retention rules.','Keep payment and external declaration actions behind explicit, contextual confirmation.']:
        doc.add_paragraph(ptxt,style='List Bullet')
    doc.add_heading('8 Production readiness checklist',level=1)
    for ptxt in ['Authentication and role policies tested','Database schema and migrations reviewed','Company ownership enforced for all records','Offline queue tested with conflict recovery','External connectors use test and production environments','Secrets stored in managed runtime configuration','Camera streams never exposed as public RTSP','Audit export reconciles with business records','Accessibility and responsive tests completed','Backups, restore test and incident runbook completed']:
        doc.add_paragraph('☐ '+ptxt)
    add_footer(doc); p=OUT/'AgriPilot_Technical_Architecture_and_User_Guide.docx'; doc.save(p); return p

if __name__=='__main__':
    ci=make_class_diagram(); fi=make_flow_diagram();
    print(checklist_doc()); print(guide_doc(ci,fi))
